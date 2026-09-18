import csv
import hashlib
import io
import json
import re
import sqlite3
from collections import Counter, defaultdict
from pathlib import Path
from zipfile import BadZipFile, ZipFile


CATEGORIES = {"MARKETING", "UTILITY", "AUTHENTICATION"}
COMPLETED = {"VERIFIED", "APPROVED"}
MAX_ROWS = 30000
MAX_BYTES = 25 * 1024 * 1024
FIELDS = ["name", "body", "header", "footer", "buttons", "meta_category",
          "requested_category", "status", "language", "created_at", "family_id"]
ALIASES = {
    "name": ["name", "templateName", "template_name"],
    "body": ["body", "text", "message", "content", "template_text"],
    "header": ["header"], "footer": ["footer"], "buttons": ["buttons"],
    "meta_category": ["meta_category", "metaTemplateCategory", "meta_category_decision"],
    "requested_category": ["requested_category", "templateCategory", "category"],
    "status": ["status", "verificationStatus", "verification_status"],
    "language": ["language", "whatsAppLanguage"],
    "created_at": ["created_at", "createdAt"], "family_id": ["family_id"],
}


def string(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value).strip()


def category(value):
    value = string(value).upper().strip()
    return value if value in CATEGORIES else "UNKNOWN"


def timestamp(value):
    return string(value.get("$date", "")) if isinstance(value, dict) else string(value)


def parse_file(content, filename):
    if len(content) > MAX_BYTES:
        raise ValueError("The file exceeds the 25 MB limit.")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".xlsx":
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_BYTES * 5:
                    raise ValueError("The expanded workbook is too large.")
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = book.active
                values = sheet.iter_rows(values_only=True)
                headers = [string(v) for v in next(values, ())]
                rows = []
                for row in values:
                    if any(v is not None for v in row):
                        rows.append(dict(zip(headers, row)))
                    if len(rows) > MAX_ROWS:
                        raise ValueError("Import at most 30,000 rows at a time.")
            finally:
                book.close()
        else:
            text = content.decode("utf-8-sig").strip()
            if not text:
                raise ValueError("The file is empty.")
            if suffix in {".csv", ".tsv"}:
                rows = list(csv.DictReader(io.StringIO(text), delimiter="\t" if suffix == ".tsv" else ","))
            elif suffix == ".jsonl":
                rows = [json.loads(line) for line in text.splitlines() if line.strip()]
            else:
                rows = json.loads(text)
                if isinstance(rows, str):
                    rows = json.loads(rows)
                if isinstance(rows, dict):
                    rows = rows.get("templates", rows.get("data", [rows]))
        if not isinstance(rows, list) or not rows:
            raise ValueError("Expected a non-empty list of template records.")
        if len(rows) > MAX_ROWS:
            raise ValueError("Import at most 30,000 rows at a time.")
        if any(not isinstance(row, dict) for row in rows):
            raise ValueError("Each template must be an object or a spreadsheet row.")
        return rows
    except (json.JSONDecodeError, UnicodeDecodeError, BadZipFile, csv.Error) as exc:
        raise ValueError("Could not read the file. Use UTF-8 JSON, JSONL, CSV, TSV, or an XLSX workbook.") from exc


def suggest_mapping(rows):
    columns = list(dict.fromkeys(str(key) for row in rows[:100] for key in row))
    return {field: next((key for key in ALIASES[field] if key in columns), "") for field in FIELDS}


def button_texts(message):
    config = message.get("interactionConfiguration") or {}
    config = config.get("buttonConfiguration") or {} if isinstance(config, dict) else {}
    buttons = []
    for key in ("qrButtons", "ctaButtons"):
        for button in config.get(key, []) or []:
            if isinstance(button, dict) and button.get("text"):
                buttons.append(string(button["text"]))
    return buttons


def record_content(record):
    return "\n".join(string(record.get(field)) for field in ("header", "body", "footer", "buttons"))


def normalized_content(text):
    text = re.sub(r"\{\{.*?\}\}", " variable ", text.lower())
    text = re.sub(r"https?://\S+|www\.\S+", " url ", text)
    text = re.sub(r"\b[\w.+-]+@[\w.-]+\.[a-z]+\b", " email ", text)
    text = re.sub(r"\d+", " number ", text)
    return " ".join(re.findall(r"\w+", text, flags=re.UNICODE))


FEATURE_NAMES = [
    "placeholder_count", "distinct_placeholders", "button_count", "has_header",
    "has_footer", "body_length", "emoji_count", "exclamation_count",
    "symbol_count", "uppercase_ratio", "number_count", "url_count",
]


def button_lines(record):
    return [line for line in string(record.get("buttons")).split("\n") if line.strip()]


def structural_features(record):
    """Signals that normalized_content deliberately erases, in FEATURE_NAMES order.

    normalized_content collapses placeholders, digits and punctuation because it
    defines family grouping; these counts recover that discarded signal without
    changing how families are formed.
    """
    body = string(record.get("body"))
    placeholders = re.findall(r"\{\{.*?\}\}", body)
    letters = [c for c in body if c.isalpha()]
    return [
        float(len(placeholders)),
        float(len(set(placeholders))),
        float(len(button_lines(record))),
        1.0 if string(record.get("header")) else 0.0,
        1.0 if string(record.get("footer")) else 0.0,
        float(len(body)),
        float(sum(1 for c in body if ord(c) >= 0x2600)),
        float(body.count("!")),
        float(len(re.findall(r"[%$₹€£]", body))),
        (sum(1 for c in letters if c.isupper()) / len(letters)) if letters else 0.0,
        float(len(re.findall(r"\d+", body))),
        float(len(re.findall(r"https?://\S+|www\.\S+", body))),
    ]


def family_key(record):
    """Grouping key over every component.

    Derived at read time on purpose: normalize_record folds `family` into the
    stored `id`, so changing the stored value would change every id, duplicate
    the dataset on re-import and orphan annotations keyed by record id.
    """
    if record.get("family_id"):
        return hashlib.sha256(string(record["family_id"]).encode()).hexdigest()[:24]
    seed = normalized_content(record_content(record)) or normalized_content(string(record.get("body")))
    return hashlib.sha256(seed.encode()).hexdigest()[:24]


def normalize_record(row, mapping, source):
    nested = isinstance(row.get("messageBody"), dict)
    if nested:
        message = row["messageBody"]
        values = {
            "external_id": string(row.get("_id")), "name": string(row.get("templateName")),
            "header": string(message.get("header")), "body": string(message.get("body")),
            "footer": string(message.get("footer")), "buttons": "\n".join(button_texts(message)),
            "meta_category": category(row.get("metaTemplateCategory")),
            "requested_category": category(message.get("templateCategory")),
            "status": string(row.get("verificationStatus")).upper() or "UNKNOWN",
            "language": string(row.get("whatsAppLanguage")),
            "created_at": timestamp(row.get("createdAt")), "updated_at": timestamp(row.get("updatedAt")),
            "format": string(message.get("type")).upper() or "UNKNOWN",
            "family_id": "", "label_source": "metaTemplateCategory",
        }
    else:
        values = {field: string(row.get(mapping.get(field, ""))) for field in FIELDS}
        values.update({
            "external_id": string(row.get("id", row.get("_id"))),
            "meta_category": category(values["meta_category"]),
            "requested_category": category(values["requested_category"]),
            "status": values["status"].upper() or "UNKNOWN",
            "created_at": timestamp(row.get(mapping.get("created_at", ""))),
            "updated_at": timestamp(row.get("updated_at")),
            "format": string(row.get("format", "TEXT")).upper(),
            "label_source": mapping.get("meta_category", "") or "not supplied",
        })
    if len(values["body"]) > 20000 or len(record_content(values)) > 40000:
        raise ValueError("A template exceeds the 40,000-character content limit.")
    # Identity fields and labels are not included in classifier features.
    content = normalized_content(record_content(values))
    family_seed = values["family_id"] or normalized_content(values["body"]) or content
    values["family"] = hashlib.sha256(family_seed.encode()).hexdigest()[:24]
    values["id"] = hashlib.sha256(json.dumps(values, sort_keys=True).encode()).hexdigest()[:32]
    values["source"] = Path(source).name
    return values


def normalize_rows(rows, mapping, source):
    records = [normalize_record(row, mapping, source) for row in rows]
    return records


def candidate(record):
    return (record["status"] in COMPLETED and record["format"] == "TEXT"
            and record["meta_category"] in {"MARKETING", "UTILITY"}
            and len(record["body"].strip()) >= 10)


def training_records(records):
    families = defaultdict(list)
    for record in records:
        if candidate(record):
            families[family_key(record)].append(record)
    conflicts = {key for key, group in families.items() if len({r["meta_category"] for r in group}) > 1}
    clean = []
    for key, group in families.items():
        if key in conflicts:
            continue
        representative = max(group, key=lambda r: r["updated_at"] or r["created_at"])
        # `family` is overridden on a copy so the stored record, and the `id` that
        # embeds its original family hash, are both left untouched.
        clean.append({**representative, "family": key})
    return clean, conflicts


def summarize(records):
    clean, conflicts = training_records(records)
    candidates = [r for r in records if candidate(r)]
    counts = Counter(family_key(r) for r in records)
    return {
        "total": len(records), "categories": dict(Counter(r["meta_category"] for r in records)),
        "requested_categories": dict(Counter(r["requested_category"] for r in records)),
        "statuses": dict(Counter(r["status"] for r in records)),
        "formats": dict(Counter(r["format"] for r in records)),
        "languages": dict(Counter(r["language"] or "UNKNOWN" for r in records)),
        "sources": dict(Counter(r["source"] for r in records)),
        "category_mismatches": sum(r["requested_category"] == "UTILITY" and r["meta_category"] == "MARKETING" for r in records),
        "families": len(counts), "duplicate_records": sum(n - 1 for n in counts.values()),
        "conflicting_families": len(conflicts),
        "conflicting_records": sum(family_key(r) in conflicts for r in candidates),
        "candidate_records": len(candidates), "training_families": len(clean),
        "training_categories": dict(Counter(r["meta_category"] for r in clean)),
        "missing_bodies": sum(not r["body"].strip() for r in records),
    }


class Store:
    def __init__(self, directory):
        self.directory = Path(directory)
        self.directory.mkdir(parents=True, exist_ok=True)
        self.path = self.directory / "templates.sqlite3"
        with self.connection() as db:
            db.execute("CREATE TABLE IF NOT EXISTS records (id TEXT PRIMARY KEY, payload TEXT NOT NULL)")

    def connection(self):
        return sqlite3.connect(self.path, timeout=30)

    def records(self):
        with self.connection() as db:
            rows = db.execute("SELECT payload FROM records ORDER BY rowid DESC").fetchall()
        return [json.loads(row[0]) for row in rows]

    def revision(self):
        with self.connection() as db:
            ids = [row[0] for row in db.execute("SELECT id FROM records ORDER BY id")]
        return hashlib.sha256("".join(ids).encode()).hexdigest()

    def import_records(self, records):
        with self.connection() as db:
            before = db.total_changes
            db.executemany("INSERT OR IGNORE INTO records (id, payload) VALUES (?, ?)",
                           [(r["id"], json.dumps(r, ensure_ascii=False)) for r in records])
            added = db.total_changes - before
        return {"added": added, "already_present": len(records) - added}
